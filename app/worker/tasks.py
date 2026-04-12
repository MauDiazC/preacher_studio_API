from app.core.celery_app import celery_app
from app.services.ai_service import ai_service
from app.core.db import supabase
from app.repository.sermon_repository import sermon_repo
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import io
import re
import asyncio
import logging

logger = logging.getLogger("celery_worker")

@celery_app.task(name="tasks.get_ai_suggestions")
def get_ai_suggestions_task(sermon_id: str, title: str, content: str, style: str = "encouraging"):
    """
    Tarea de Celery para obtener sugerencias de IA para un sermón.
    """
    logger.info(f"🚀 Iniciando tarea de sugerencias para: {title}")
    try:
        result = asyncio.run(ai_service.get_suggestions(title, content, style))
        
        # Guardar snapshot de historial
        suggestion_str = str(result.model_dump())
        sermon_repo.save_history_snapshot(supabase, sermon_id, suggestion_str, "AI_LOG")
        
        return result.model_dump()
    except Exception as e:
        logger.error(f"❌ Error en tarea get_ai_suggestions: {e}")
        return {"error": str(e)}

@celery_app.task(name="tasks.analyze_verse_exegesis")
def analyze_verse_exegesis_task(verse_reference: str, language: str = "es"):
    """
    Tarea de Celery para realizar un análisis exegético de un versículo.
    """
    logger.info(f"🚀 Iniciando tarea de exégesis para: {verse_reference}")
    try:
        result = asyncio.run(ai_service.analyze_verse(verse_reference, language))
        return result.model_dump()
    except Exception as e:
        logger.error(f"❌ Error en tarea analyze_verse_exegesis: {e}")
        return {"error": str(e)}

@celery_app.task(name="tasks.export_sermon_pdf")
def export_sermon_pdf_task(sermon_id: str, user_id: str):
    """
    Genera un PDF y lo sube a Supabase Storage.
    """
    logger.info(f"🚀 Exportando PDF para sermón: {sermon_id}")
    try:
        res = sermon_repo.get_by_id(supabase, sermon_id, user_id)
        if not res.data:
            return {"error": "Sermón no encontrado"}

        sermon = res.data
        buffer = io.BytesIO()
        p = canvas.Canvas(buffer, pagesize=letter)

        p.setFont("Helvetica-Bold", 18)
        p.drawString(50, 750, sermon['title'])
        p.setFont("Helvetica-Oblique", 12)
        p.drawString(50, 730, f"Pasaje: {sermon.get('main_passage') or 'N/A'}")
        p.setFont("Helvetica", 11)
        text_object = p.beginText(50, 700)
        content = sermon.get("content", "")
        for line in content.split("\n"):
            text_object.textLine(line)
        p.drawText(text_object)
        p.showPage()
        p.save()
        buffer.seek(0)

        # Subir a Supabase
        file_name = f"pdf/{user_id}/{sermon_id}.pdf"
        supabase.storage.from_("temporary_exports").upload(
            file_name, 
            buffer.getvalue(),
            {"content-type": "application/pdf", "upsert": "true"}
        )

        # Generar URL firmada (60 minutos)
        signed_url = supabase.storage.from_("temporary_exports").create_signed_url(file_name, 3600)
        return {"download_url": signed_url["signedURL"]}
    except Exception as e:
        logger.error(f"❌ Error exportando PDF: {e}")
        return {"error": str(e)}

@celery_app.task(name="tasks.export_sermon_pptx")
def export_sermon_pptx_task(sermon_id: str, user_id: str):
    """
    Genera un PPTX y lo sube a Supabase Storage.
    """
    logger.info(f"🚀 Exportando PPTX para sermón: {sermon_id}")
    try:
        res = sermon_repo.get_by_id(supabase, sermon_id, user_id)
        if not res.data:
            return {"error": "Sermón no encontrado"}

        sermon = res.data
        prs = Presentation()

        BG_COLOR = RGBColor(13, 11, 31)
        ACCENT_PURPLE = RGBColor(125, 92, 255)
        ACCENT_GOLD = RGBColor(255, 184, 120)
        TEXT_MAIN = RGBColor(255, 255, 255)
        TEXT_DIM = RGBColor(161, 161, 170)

        def apply_sacred_style(slide, is_title=False):
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = BG_COLOR
            if not is_title:
                txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(9), Inches(0.3))
                tf = txBox.text_frame
                p = tf.paragraphs[0]
                p.text = "PREACHER STUDIO | Asistencia Homilética Digital"
                p.font.size = Pt(9)
                p.font.color.rgb = TEXT_DIM
                p.alignment = PP_ALIGN.LEFT

        slide = prs.slides.add_slide(prs.slide_layouts[0])
        apply_sacred_style(slide, is_title=True)
        slide.shapes.title.text = sermon["title"].upper()
        slide.placeholders[1].text = f"ANÁLISIS EXEGÉTICO & HOMILÉTICO\n{sermon.get('main_passage') or 'Estudio Bíblico'}"

        content = sermon.get("content", "")
        sections = re.split(r'(\d+\.\s+[A-ZÁÉÍÓÚÑ\s\(\)]+:|VERSIÓN [A-Z0-9\s]+:)', content)
        
        slides_data = []
        for i in range(1, len(sections), 2):
            header = sections[i].strip()
            body = sections[i+1].strip() if i+1 < len(sections) else ""
            if body: slides_data.append((header, body))

        bullet_layout = prs.slide_layouts[1]
        for s_title, s_body in slides_data:
            chunks = [s_body[i:i+650] for i in range(0, len(s_body), 650)]
            for idx, chunk in enumerate(chunks):
                slide = prs.slides.add_slide(bullet_layout)
                apply_sacred_style(slide)
                slide.shapes.title.text = s_title + (" (cont.)" if idx > 0 else "")
                slide.placeholders[1].text = chunk

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)

        # Subir a Supabase
        file_name = f"pptx/{user_id}/{sermon_id}.pptx"
        supabase.storage.from_("temporary_exports").upload(
            file_name, 
            buffer.getvalue(),
            {"content-type": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "upsert": "true"}
        )

        signed_url = supabase.storage.from_("temporary_exports").create_signed_url(file_name, 3600)
        return {"download_url": signed_url["signedURL"]}
    except Exception as e:
        logger.error(f"❌ Error exportando PPTX: {e}")
        return {"error": str(e)}
